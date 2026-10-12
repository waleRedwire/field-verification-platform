"""
FIELD VERIFICATION PLATFORM
Per-user projects · Persistent rules · Admin Analysis/Users · Security
"""
import streamlit as st
import pandas as pd
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils.dataframe import dataframe_to_rows

from db import (
    create_project, list_projects, get_project,
    get_rules, save_rules, save_dataset, save_run,
    create_user, list_users, set_user_active,
    user_can_access_project,
)
from auth_ui import require_login, logout_button
from verification import run_full_verification
from map_utils import create_data_map, find_nearest_row
from mesbi import run_mesbi
from audio import download_audio_bytes, transcribe_english
from security_utils import (
    escape_text,
    sanitize_dataframe_for_excel,
    check_upload_limits,
    safe_error_message,
    MAX_ROWS,
)

try:
    from google_utils import load_google_sheet
    GOOGLE_AVAILABLE = True
except ImportError:
    GOOGLE_AVAILABLE = False

st.set_page_config(
    page_title="FIELD VERIFICATION PLATFORM",
    page_icon="✅",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""
<style>
    .stApp {
        background: linear-gradient(180deg, #f1f5f9 0%, #f8fafc 40%, #ffffff 100%);
        color: #0f172a;
    }
    .main .block-container {
        padding-top: 1.25rem;
        padding-bottom: 2.5rem;
        max-width: 1280px;
    }
    .app-header {
        background: linear-gradient(135deg, #1e3a5f 0%, #1e40af 55%, #2563eb 100%);
        border-radius: 16px;
        padding: 1.35rem 1.75rem;
        margin-bottom: 1.25rem;
        color: white;
        box-shadow: 0 10px 30px rgba(30, 64, 175, 0.18);
    }
    .app-header h1 {
        margin: 0;
        font-size: 1.45rem;
        font-weight: 700;
        color: #ffffff;
    }
    .app-header .subtitle {
        color: rgba(219, 234, 254, 0.9);
        font-size: 0.9rem;
        margin-top: 0.35rem;
    }
    section[data-testid="stSidebar"] {
        background: #ffffff !important;
        border-right: 1px solid #e2e8f0;
    }
    div[data-testid="stMetric"] {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 14px;
        padding: 1rem 1.15rem;
    }
    .stButton > button[kind="primary"] {
        background: linear-gradient(135deg, #2563eb, #1d4ed8) !important;
        border: none !important;
        color: white !important;
        border-radius: 10px !important;
        font-weight: 600 !important;
    }
    .stTabs [data-baseweb="tab-list"] {
        gap: 4px;
        background: #ffffff;
        padding: 6px;
        border-radius: 12px;
        border: 1px solid #e2e8f0;
    }
    .stTabs [data-baseweb="tab"] {
        border-radius: 8px;
        color: #64748b !important;
        font-weight: 500;
        padding: 8px 14px;
    }
    .stTabs [aria-selected="true"] {
        background: #eff6ff !important;
        color: #1e40af !important;
        font-weight: 600 !important;
        border: 1px solid #bfdbfe !important;
    }
    #MainMenu { visibility: hidden; }
    footer { visibility: hidden; }
</style>
""", unsafe_allow_html=True)

user = require_login()
is_admin = user.get("role") == "admin"

# ---------- Session state ----------
defaults = {
    "current_project_id": None,
    "df": None,
    "last_report": None,
    "demo_by_state": {},
    "mesbi_result": None,
    "selected_point_idx": None,
    "point_audio_bytes": None,
    "point_audio_format": None,
    "point_audio_url": None,
    "point_transcript": None,
}
for k, v in defaults.items():
    if k not in st.session_state:
        st.session_state[k] = v


def clear_working_data():
    st.session_state.df = None
    st.session_state.last_report = None
    st.session_state.mesbi_result = None
    st.session_state.selected_point_idx = None
    st.session_state.point_audio_bytes = None
    st.session_state.point_transcript = None
    st.session_state.point_audio_url = None
    st.session_state.demo_by_state = {}


def load_data(uploaded_file) -> pd.DataFrame:
    if uploaded_file.name.endswith((".xlsx", ".xls")):
        return pd.read_excel(uploaded_file)
    return pd.read_csv(uploaded_file)


def build_colour_coded_excel(pass_fail_df: pd.DataFrame) -> bytes:
    safe_df = sanitize_dataframe_for_excel(pass_fail_df)
    wb = Workbook()
    ws = wb.active
    ws.title = "Verification Results"
    for r_idx, row in enumerate(dataframe_to_rows(safe_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if r_idx > 1 and c_idx == 1:
                if str(value).upper() == "FAIL":
                    cell.fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
                    cell.font = Font(bold=True, color="FFFFFF")
                else:
                    cell.fill = PatternFill(start_color="6BCB77", end_color="6BCB77", fill_type="solid")
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def load_and_transcribe_point(audio_url: str, use_api: bool, api_key, model_size: str):
    data, suffix, err = download_audio_bytes(audio_url)
    if err:
        st.session_state.point_audio_bytes = None
        st.session_state.point_transcript = {"error": err, "text": None}
        return
    st.session_state.point_audio_bytes = data
    st.session_state.point_audio_format = (suffix or ".mp3").lstrip(".")
    st.session_state.point_audio_url = audio_url
    out = transcribe_english(
        audio_url,
        model_size=model_size,
        use_openai_api=use_api,
        openai_api_key=api_key
    )
    st.session_state.point_transcript = out


# ---------- Sidebar ----------
st.sidebar.markdown("#### Account")
st.sidebar.write(f"**{escape_text(user['name'])}**")
st.sidebar.caption(f"{escape_text(user['username'])} · {escape_text(user['role'])}")
logout_button()
st.sidebar.markdown("---")
st.sidebar.markdown("#### Your projects")
st.sidebar.caption("Only your projects are listed" if not is_admin else "Admin: all projects visible")

projects = list_projects(user)
project_names = {p["name"]: p["id"] for p in projects}

selected_name = st.sidebar.selectbox(
    "Select project",
    options=["— create new —"] + list(project_names.keys()),
    label_visibility="collapsed"
)

if selected_name == "— create new —":
    with st.sidebar.form("new_project"):
        new_name = st.text_input("Project name")
        new_desc = st.text_area("Description", height=80)
        if st.form_submit_button("Create Project"):
            if new_name.strip():
                try:
                    pid = create_project(
                        new_name.strip(),
                        new_desc.strip(),
                        owner_id=user["id"]
                    )
                    clear_working_data()
                    st.session_state.current_project_id = pid
                    st.rerun()
                except Exception:
                    st.sidebar.error("Could not create project (name may already exist for you).")
            else:
                st.sidebar.error("Name required")
else:
    new_id = project_names[selected_name]
    if st.session_state.current_project_id != new_id:
        clear_working_data()
        st.session_state.current_project_id = new_id

st.sidebar.markdown("---")
st.sidebar.caption("Field Verification Platform")

# ---------- Main shell ----------
if st.session_state.current_project_id is None:
    st.markdown("""
    <div class="app-header"><div>
        <h1>FIELD VERIFICATION PLATFORM</h1>
        <div class="subtitle">Create or select a project in the sidebar to begin</div>
    </div></div>
    """, unsafe_allow_html=True)
    st.info("Your projects and verification rules are saved to your account and will still be here next login.")
    if not is_admin:
        st.stop()

project = None
rules = {}
if st.session_state.current_project_id is not None:
    if not user_can_access_project(st.session_state.current_project_id, user):
        st.error("You do not have access to this project.")
        st.session_state.current_project_id = None
        st.stop()
    project = get_project(st.session_state.current_project_id, user)
    if project is None:
        st.error("Project not found or access denied.")
        st.session_state.current_project_id = None
        st.stop()
    rules = get_rules(st.session_state.current_project_id, user)
    if not st.session_state.demo_by_state and rules.get("demographic_by_state"):
        st.session_state.demo_by_state = rules.get("demographic_by_state", {})
    safe_proj = escape_text(project["name"])
    safe_desc = escape_text(project.get("description") or "")
    st.markdown(f"""
    <div class="app-header"><div>
        <h1>FIELD VERIFICATION PLATFORM</h1>
        <div class="subtitle">{safe_proj}{(' — ' + safe_desc) if safe_desc else ''}</div>
    </div></div>
    """, unsafe_allow_html=True)

# Tabs
if is_admin:
    if st.session_state.current_project_id is not None:
        tab_upload, tab_rules, tab_verify, tab_map, tab_report, tab_analysis, tab_users = st.tabs([
            "1. Upload Data",
            "2. Verification Rules",
            "3. Run Verification",
            "4. Map",
            "5. Report",
            "6. Analysis",
            "7. Users"
        ])
    else:
        tab_users = st.tabs(["7. Users"])[0]
        tab_upload = tab_rules = tab_verify = tab_map = tab_report = tab_analysis = None
else:
    tab_upload, tab_rules, tab_verify, tab_map, tab_report = st.tabs([
        "1. Upload Data",
        "2. Verification Rules",
        "3. Run Verification",
        "4. Map",
        "5. Report"
    ])
    tab_analysis = None
    tab_users = None

# ========== USERS (admin) ==========
if tab_users is not None:
    with tab_users:
        st.subheader("User management")
        st.caption("Each user only sees their own projects. Rules persist per project.")

        with st.form("create_user_form"):
            c1, c2 = st.columns(2)
            with c1:
                new_username = st.text_input("Username *")
                new_name = st.text_input("Full name *")
            with c2:
                new_email = st.text_input("Email *")
                new_phone = st.text_input("Phone number")
            new_role = st.selectbox("Role", ["user", "admin"])
            create_btn = st.form_submit_button("Create user", type="primary")

        if create_btn:
            result = create_user(new_username, new_email, new_name, new_phone, role=new_role)
            if not result["ok"]:
                st.error(result["error"])
            else:
                st.success(f"User **{escape_text(new_username)}** created.")
                st.warning("Copy the default password now — it will not be shown again.")
                st.code(result["default_password"], language=None)

        st.markdown("### Existing users")
        users = list_users()
        if not users:
            st.info("No users yet.")
        else:
            for u in users:
                cols = st.columns([2, 2, 2, 1, 1, 1])
                cols[0].write(f"**{escape_text(u['username'])}**")
                cols[1].write(escape_text(u["email"]))
                cols[2].write(escape_text(u.get("name") or ""))
                cols[3].write(escape_text(u["role"]))
                cols[4].write("Active" if u["is_active"] else "Off")
                if u["id"] != user["id"]:
                    if u["is_active"]:
                        if cols[5].button("Disable", key=f"dis_{u['id']}"):
                            set_user_active(u["id"], False)
                            st.rerun()
                    else:
                        if cols[5].button("Enable", key=f"en_{u['id']}"):
                            set_user_active(u["id"], True)
                            st.rerun()

if st.session_state.current_project_id is None or tab_upload is None:
    st.stop()

# ========== 1. UPLOAD ==========
with tab_upload:
    st.subheader("Load Data")
    st.caption(f"Limits: max file 25 MB · max {MAX_ROWS:,} rows")
    source = st.radio("Data source", options=["Upload Excel / CSV", "Google Sheet"], horizontal=True)
    if source == "Upload Excel / CSV":
        uploaded = st.file_uploader("Choose file", type=["csv", "xlsx", "xls"])
        if uploaded:
            err = check_upload_limits(uploaded)
            if err:
                st.error(err)
            else:
                try:
                    df = load_data(uploaded)
                    err2 = check_upload_limits(None, df)
                    if err2:
                        st.error(err2)
                    else:
                        st.session_state.df = df
                        st.success(f"Loaded **{len(df):,}** rows × **{len(df.columns)}** columns")
                        st.dataframe(df.head(20), use_container_width=True)
                        if st.button("Save dataset metadata to project", type="secondary"):
                            ds_id = save_dataset(
                                st.session_state.current_project_id,
                                uploaded.name,
                                df,
                                user=user
                            )
                            if ds_id:
                                st.success(f"Dataset metadata saved (id={ds_id})")
                            else:
                                st.error("Could not save (access denied).")
                except Exception:
                    st.error(safe_error_message(Exception()))
    else:
        if not GOOGLE_AVAILABLE:
            st.error("Google libraries not installed.")
        else:
            sa_file = st.text_input("Service Account JSON path", value="service_account.json")
            sheet_url = st.text_input("Google Sheet URL or ID")
            worksheet_name = st.text_input("Worksheet name (blank = first)", value="")
            if st.button("Load from Google Sheet", type="primary"):
                try:
                    with st.spinner("Loading..."):
                        df = load_google_sheet(sa_file, sheet_url.strip(), worksheet_name.strip() or None)
                    err2 = check_upload_limits(None, df)
                    if err2:
                        st.error(err2)
                    else:
                        st.session_state.df = df
                        st.success(f"Loaded **{len(df):,}** rows")
                        st.dataframe(df.head(20), use_container_width=True)
                except Exception:
                    st.error(safe_error_message(Exception()))
    if st.session_state.df is not None:
        st.caption("Columns: " + ", ".join(escape_text(c) for c in st.session_state.df.columns))

# ========== 2. RULES ==========
with tab_rules:
    st.subheader("Verification Settings")
    st.caption("Saved to this project and restored when you log in again.")
    all_cols = list(st.session_state.df.columns) if st.session_state.df is not None else []

    with st.expander("Geographic and Agent Columns", expanded=True):
        c1, c2 = st.columns(2)
        with c1:
            lat_col = st.selectbox(
                "Latitude column", options=[None] + all_cols,
                index=([None] + all_cols).index(rules.get("lat_col")) if rules.get("lat_col") in ([None] + all_cols) else 0
            )
            agent_col = st.selectbox(
                "Agent / Enumerator column", options=[None] + all_cols,
                index=([None] + all_cols).index(rules.get("agent_col")) if rules.get("agent_col") in ([None] + all_cols) else 0
            )
            audio_col = st.selectbox(
                "Audio / Interview URL column", options=[None] + all_cols,
                index=([None] + all_cols).index(rules.get("audio_col")) if rules.get("audio_col") in ([None] + all_cols) else 0
            )
        with c2:
            lon_col = st.selectbox(
                "Longitude column", options=[None] + all_cols,
                index=([None] + all_cols).index(rules.get("lon_col")) if rules.get("lon_col") in ([None] + all_cols) else 0
            )
            state_col = st.selectbox(
                "State column", options=[None] + all_cols,
                index=([None] + all_cols).index(rules.get("state_col")) if rules.get("state_col") in ([None] + all_cols) else 0
            )
            submission_date_col = st.selectbox(
                "Submission / interview date column",
                options=[None] + all_cols,
                index=([None] + all_cols).index(rules.get("submission_date_col"))
                if rules.get("submission_date_col") in ([None] + all_cols) else 0
            )

    with st.expander("Duplicate Detection"):
        key_cols = st.multiselect(
            "Exact-duplicate key columns", options=all_cols,
            default=[c for c in rules.get("duplicate_keys", []) if c in all_cols]
        )
        fuzzy_thresh = st.slider("Fuzzy match threshold", 70, 100, rules.get("fuzzy_threshold", 90))

    with st.expander("Interview Duration Settings"):
        c1, c2 = st.columns(2)
        with c1:
            start_col = st.selectbox(
                "Start Time column", options=[None] + all_cols,
                index=([None] + all_cols).index(rules.get("start_time_col")) if rules.get("start_time_col") in ([None] + all_cols) else 0
            )
            min_duration = st.number_input(
                "Minimum duration (minutes)", 0.5, 180.0,
                float(rules.get("min_duration_minutes", 5)), 0.5
            )
        with c2:
            end_col = st.selectbox(
                "End Time column", options=[None] + all_cols,
                index=([None] + all_cols).index(rules.get("end_time_col")) if rules.get("end_time_col") in ([None] + all_cols) else 0
            )
            max_duration = st.number_input(
                "Maximum duration (minutes)", 1.0, 480.0,
                float(rules.get("max_duration_minutes", 120)), 1.0
            )

    with st.expander("Geographic Cluster Settings"):
        cluster_enabled = st.checkbox("Enable cluster detection", value=rules.get("cluster_enabled", True))
        c1, c2 = st.columns(2)
        with c1:
            min_points = st.number_input(
                "Minimum points to form a cluster", 2, 50,
                int(rules.get("cluster_min_points", 5)), 1
            )
        with c2:
            radius_m = st.number_input(
                "Radius (meters)", 10, 5000,
                int(rules.get("cluster_radius_meters", 100)), 10
            )

    with st.expander("Agent Pattern / Fraud Thresholds"):
        min_ent = st.number_input(
            "Minimum answer entropy", 0.5, 5.0,
            rules.get("pattern_checks", {}).get("min_entropy", 1.5), 0.1
        )
        max_same = st.slider(
            "Max same-answer percentage", 0.3, 1.0,
            rules.get("pattern_checks", {}).get("max_same_answer_pct", 0.6), 0.05
        )
        short_sec = st.number_input(
            "Flag fill interval shorter than (seconds)", 5, 300,
            rules.get("pattern_checks", {}).get("flag_short_fill_time_seconds", 30)
        )
        ts_col = st.selectbox(
            "Timestamp column (fill-speed)", options=[None] + all_cols,
            index=([None] + all_cols).index(rules.get("timestamp_col")) if rules.get("timestamp_col") in ([None] + all_cols) else 0
        )

    with st.expander("Demographic Targets – Per State", expanded=True):
        c1, c2 = st.columns(2)
        with c1:
            gender_col = st.selectbox(
                "Gender column", options=[None] + all_cols,
                index=([None] + all_cols).index(rules.get("gender_col")) if rules.get("gender_col") in ([None] + all_cols) else 0,
                key="gcol"
            )
        with c2:
            age_col = st.selectbox(
                "Age column", options=[None] + all_cols,
                index=([None] + all_cols).index(rules.get("age_col")) if rules.get("age_col") in ([None] + all_cols) else 0,
                key="acol"
            )
        if st.session_state.df is not None and state_col:
            states = sorted(st.session_state.df[state_col].dropna().astype(str).unique().tolist())
            age_values = (
                sorted(st.session_state.df[age_col].dropna().astype(str).unique().tolist())
                if age_col and age_col in st.session_state.df.columns else []
            )
            for state in states:
                with st.expander(f"State: {escape_text(state)}", expanded=False):
                    current = st.session_state.demo_by_state.get(state, {})
                    cg, ca = current.get("gender", {}), current.get("age", {})
                    g1, g2 = st.columns(2)
                    with g1:
                        male = st.number_input(
                            f"{state} Male", 0, 100000, int(cg.get("Male", 0)), 1, key=f"m_{state}"
                        )
                    with g2:
                        female = st.number_input(
                            f"{state} Female", 0, 100000, int(cg.get("Female", 0)), 1, key=f"f_{state}"
                        )
                    age_targets = {}
                    if age_values:
                        cols_a = st.columns(min(4, len(age_values)))
                        for i, val in enumerate(age_values):
                            with cols_a[i % len(cols_a)]:
                                age_targets[val] = st.number_input(
                                    f"{state} {val}", 0, 100000, int(ca.get(val, 0)), 1, key=f"a_{state}_{val}"
                                )
                    else:
                        age_targets = ca
                    st.session_state.demo_by_state[state] = {
                        "gender": {"Male": male, "Female": female},
                        "age": age_targets
                    }

    if st.button("Save All Rules", type="primary"):
        new_rules = {
            "duplicate_keys": key_cols,
            "fuzzy_threshold": fuzzy_thresh,
            "lat_col": lat_col,
            "lon_col": lon_col,
            "agent_col": agent_col,
            "audio_col": audio_col,
            "submission_date_col": submission_date_col,
            "state_col": state_col,
            "timestamp_col": ts_col,
            "start_time_col": start_col,
            "end_time_col": end_col,
            "min_duration_minutes": min_duration,
            "max_duration_minutes": max_duration,
            "cluster_enabled": cluster_enabled,
            "cluster_min_points": min_points,
            "cluster_radius_meters": radius_m,
            "pattern_checks": {
                "min_entropy": min_ent,
                "max_same_answer_pct": max_same,
                "flag_short_fill_time_seconds": short_sec
            },
            "gender_col": gender_col,
            "age_col": age_col,
            "demographic_targets": {"gender": {}, "age": {}},
            "demographic_by_state": st.session_state.get("demo_by_state", {})
        }
        ok = save_rules(st.session_state.current_project_id, new_rules, user=user)
        if ok:
            st.success("Rules saved. They will still be here next time you log in.")
            st.rerun()
        else:
            st.error("Could not save rules (access denied).")

# ========== 3. VERIFY ==========
with tab_verify:
    st.subheader("Run Full Verification")
    if st.session_state.df is None:
        st.warning("Upload data first.")
    else:
        if len(st.session_state.df) > MAX_ROWS:
            st.error(f"Dataset exceeds {MAX_ROWS:,} rows.")
        elif st.button("Run Full Verification", type="primary"):
            try:
                with st.spinner("Running..."):
                    report = run_full_verification(st.session_state.df, rules)
                    st.session_state.last_report = report
                    save_run(st.session_state.current_project_id, None, report, user=user)
                st.success("Done — open Report tab.")
            except Exception:
                st.error(safe_error_message(Exception()))
        if st.session_state.last_report:
            r = st.session_state.last_report
            c1, c2, c3, c4, c5 = st.columns(5)
            c1.metric("Rows", r["n_rows"])
            c2.metric("Duplicates", r["duplicates"]["exact"]["count"])
            c3.metric("Flagged Agents", r.get("patterns", {}).get("summary", {}).get("n_flagged", 0))
            c4.metric("Short Interviews", r.get("duration_check", {}).get("flagged_short_count", 0))
            c5.metric("Clusters", r.get("cluster_check", {}).get("flagged_count", 0))

# ========== 4. MAP ==========
with tab_map:
    st.subheader("Geographic Map")
    st.caption("Click a point for Agent, date, audio and English transcript.")

    if st.session_state.df is None:
        st.warning("Upload data first.")
    else:
        lat = rules.get("lat_col")
        lon = rules.get("lon_col")
        audio_col = rules.get("audio_col")
        agent_col = rules.get("agent_col")
        date_col = rules.get("submission_date_col") or rules.get("timestamp_col")
        state_col = rules.get("state_col")

        if not lat or not lon:
            st.info("Set Latitude and Longitude in Verification Rules.")
        elif lat not in st.session_state.df.columns or lon not in st.session_state.df.columns:
            st.error("Lat/lon columns not found. Re-save rules after uploading data.")
        else:
            df_map = st.session_state.df.copy()
            failed_indexes = set()
            if st.session_state.last_report:
                failed_indexes = set(st.session_state.last_report.get("failed_indexes", []))

            f1, f2, f3 = st.columns(3)
            with f1:
                if agent_col and agent_col in df_map.columns:
                    agents = sorted(df_map[agent_col].dropna().astype(str).unique().tolist())
                    sel_a = st.multiselect("Filter Agent", agents, default=agents)
                    if sel_a:
                        df_map = df_map[df_map[agent_col].astype(str).isin(sel_a)]
            with f2:
                if state_col and state_col in df_map.columns:
                    states = sorted(df_map[state_col].dropna().astype(str).unique().tolist())
                    sel_s = st.multiselect("Filter State", states, default=states)
                    if sel_s:
                        df_map = df_map[df_map[state_col].astype(str).isin(sel_s)]
            with f3:
                heat = st.checkbox("Heatmap")
                satellite = st.checkbox("Satellite")
                use_api = st.checkbox("Whisper via OpenAI API", value=False)
                model_size = st.selectbox("Local model", ["base", "tiny", "small"], index=0)

            api_key = None
            if use_api:
                try:
                    api_key = st.secrets.get("OPENAI_API_KEY")
                except Exception:
                    api_key = None
                if not api_key:
                    api_key = st.text_input("OpenAI API key", type="password")

            st.caption(f"**{len(df_map)}** points · Red = failed · Green = clean")

            if len(df_map) > 0:
                m = create_data_map(
                    df_map, lat_col=lat, lon_col=lon,
                    agent_col=agent_col, state_col=state_col,
                    audio_col=audio_col, date_col=date_col,
                    heat=heat, satellite=satellite,
                    failed_indexes=failed_indexes
                )
                try:
                    from streamlit_folium import st_folium
                    map_state = st_folium(
                        m, width=1200, height=600,
                        returned_objects=["last_object_clicked", "last_clicked"]
                    )
                except ImportError:
                    st.components.v1.html(m._repr_html_(), height=600)
                    map_state = None

                click_lat = click_lon = None
                if map_state:
                    obj = map_state.get("last_object_clicked") or {}
                    if obj.get("lat") is not None:
                        click_lat = obj.get("lat")
                        click_lon = obj.get("lng") or obj.get("lon")
                    elif map_state.get("last_clicked"):
                        click_lat = map_state["last_clicked"].get("lat")
                        click_lon = map_state["last_clicked"].get("lng")

                if click_lat is not None and click_lon is not None:
                    nearest = find_nearest_row(df_map, lat, lon, float(click_lat), float(click_lon))
                    if nearest is not None and nearest != st.session_state.selected_point_idx:
                        st.session_state.selected_point_idx = nearest
                        st.session_state.point_audio_bytes = None
                        st.session_state.point_transcript = None
                        st.session_state.point_audio_url = None
                        row = df_map.loc[nearest]
                        if audio_col and audio_col in df_map.columns:
                            url = str(row.get(audio_col, "")).strip()
                            if url and url.lower() not in ("nan", "none", "") and url.lower().startswith(("http://", "https://")):
                                with st.spinner("Loading audio & transcribing..."):
                                    load_and_transcribe_point(url, use_api, api_key, model_size)
                        st.rerun()

            idx = st.session_state.selected_point_idx
            if idx is not None and idx in st.session_state.df.index:
                row = st.session_state.df.loc[idx]
                st.markdown("#### Selected interview")
                d1, d2, d3 = st.columns(3)
                with d1:
                    agent_val = row[agent_col] if agent_col and agent_col in row.index else "—"
                    st.markdown(f"**Agent**  \n{escape_text(agent_val)}")
                with d2:
                    date_val = row[date_col] if date_col and date_col in row.index else "—"
                    st.markdown(f"**Submission date**  \n{escape_text(date_val)}")
                with d3:
                    state_val = row[state_col] if state_col and state_col in row.index else "—"
                    st.markdown(f"**State**  \n{escape_text(state_val)}")

                if audio_col and audio_col in row.index:
                    url = str(row.get(audio_col, "")).strip()
                    if url and url.lower() not in ("nan", "none", ""):
                        if url.lower().startswith(("http://", "https://")):
                            st.markdown(f"**Interview** · [source link]({url})")
                        if st.session_state.point_audio_bytes and st.session_state.point_audio_url == url:
                            fmt = st.session_state.point_audio_format or "mp3"
                            st.audio(st.session_state.point_audio_bytes, format=f"audio/{fmt}")
                        if st.session_state.point_transcript:
                            tr = st.session_state.point_transcript
                            if tr.get("error"):
                                st.error(escape_text(tr["error"]))
                            else:
                                st.markdown("**English transcript**")
                                st.text_area(
                                    "Transcript",
                                    value=tr.get("text") or "",
                                    height=200,
                                    label_visibility="collapsed"
                                )
            else:
                st.info("Click a point on the map to open the interview.")

# ========== 5. REPORT ==========
with tab_report:
    st.subheader("Verification Report")
    if not st.session_state.last_report:
        st.info("Run verification first.")
    else:
        r = st.session_state.last_report
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Duplicates", r["duplicates"]["exact"]["count"])
        c2.metric("Short Interviews", r.get("duration_check", {}).get("flagged_short_count", 0))
        c3.metric("Clusters", r.get("cluster_check", {}).get("flagged_count", 0))
        c4.metric("Suspicious Agents", r.get("patterns", {}).get("summary", {}).get("n_flagged", 0))
        pass_fail_df = r.get("pass_fail_table")
        if pass_fail_df is not None:
            st.dataframe(pass_fail_df, use_container_width=True)
        demo = r.get("demographics", {})
        if demo.get("gender_table"):
            st.markdown("#### Gender by State")
            st.dataframe(pd.DataFrame(demo["gender_table"]), use_container_width=True)
        if demo.get("age_table"):
            st.markdown("#### Age by State")
            st.dataframe(pd.DataFrame(demo["age_table"]), use_container_width=True)
        if pass_fail_df is not None:
            st.download_button(
                "Download Excel",
                data=build_colour_coded_excel(pass_fail_df),
                file_name=f"verification_report_{project['name']}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

# ========== 6. ANALYSIS (admin) ==========
if tab_analysis is not None:
    with tab_analysis:
        st.subheader("Analysis")
        st.caption("Admin only")
        if st.session_state.df is None:
            st.warning("Upload data first.")
        else:
            if st.button("Run MESBI", type="primary"):
                try:
                    with st.spinner("MESBI..."):
                        st.session_state.mesbi_result = run_mesbi(st.session_state.df, match_threshold=90)
                except Exception:
                    st.error(safe_error_message(Exception()))
            if st.session_state.mesbi_result:
                r = st.session_state.mesbi_result
                if r.get("mesbi_score") is not None:
                    st.metric("MESBI Score", f"{r['mesbi_score']:.4f}")
                if r.get("indicator_averages"):
                    st.dataframe(pd.DataFrame([
                        {"Indicator": k, "Average": v} for k, v in r["indicator_averages"].items()
                    ]), use_container_width=True)
                if r.get("dimension_scores"):
                    st.dataframe(pd.DataFrame([
                        {
                            "Dimension": k,
                            "Name": v["name"],
                            "Score": v["score"],
                            "Weight": v["weight"]
                        }
                        for k, v in r["dimension_scores"].items()
                    ]), use_container_width=True)